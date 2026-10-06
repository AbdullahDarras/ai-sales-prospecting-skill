import openpyxl
import pytest

from sales_agent.demo_seed import seed_demo
from sales_agent.export_xlsx import SHEET_REVIEW, SHEET_EXCLUDED, SHEET_NOTES, SHEET_SUMMARY, export_xlsx
from sales_agent.states import State
from sales_agent.store import Store


@pytest.fixture
def store(tmp_path):
    return Store(tmp_path / "leads.db")


def lead(store, name, code="VIS-3", country="SA", city="الرياض", path=(), verification=None, fit=None,
         fields=None, concerns=(), reason="سبب"):
    lead_id = store.add_lead(code, country, city, name)
    for state in path:
        store.transition(lead_id, state, reason)
    if verification is not None:
        store.set_field(lead_id, "verification_score", str(verification), "system", "system", 1.0)
    if fit is not None:
        store.set_field(lead_id, "fit_score", str(fit), "system", "system", 1.0)
    for k, (v, url, cls) in (fields or {}).items():
        store.set_field(lead_id, k, v, url, cls, 0.8)
    for c in concerns:
        store.set_field(lead_id, "concern", c, "https://d2020.example/x", "web", 0.7)
    return lead_id


UNDECIDED = (State.ENRICHED, State.VERIFIED, State.SCORED, State.UNDECIDED)
OUT = (State.ENRICHED, State.VERIFIED, State.SCORED, State.OUT_OF_ICP)
DISCARDED = (State.ENRICHED, State.DISCARDED)


def build(store, tmp_path):
    out = tmp_path / "leads.xlsx"
    export_xlsx(store, out)
    return openpyxl.load_workbook(out)


def rows(ws):
    return [[c.value for c in r] for r in ws.iter_rows(min_row=2)]


def headers(ws):
    return [c.value for c in ws[1]]


def test_workbook_has_expected_sheets_in_order(store, tmp_path):
    wb = build(store, tmp_path)
    assert wb.sheetnames == [SHEET_SUMMARY, SHEET_REVIEW, SHEET_EXCLUDED, SHEET_NOTES]


def test_review_sheet_contains_only_viable_leads_sorted_by_verification(store, tmp_path):
    lead(store, "منخفض", path=UNDECIDED, verification=55, fit=30)
    lead(store, "مرتفع", path=UNDECIDED, verification=63, fit=55)
    lead(store, "خارج", path=OUT, verification=84, fit=0)
    ws = build(store, tmp_path)[SHEET_REVIEW]
    names = [r[headers(ws).index("الاسم")] for r in rows(ws)]
    assert names == ["مرتفع", "منخفض"]


def test_excluded_sheet_has_out_of_icp_and_discarded_with_reason(store, tmp_path):
    lead(store, "كبيرة", path=OUT, verification=84, reason="أكثر من 50 موظفًا (103)")
    lead(store, "ضعيفة", path=DISCARDED, verification=25, reason="تحقق منخفض (25)")
    lead(store, "للمراجعة", path=UNDECIDED, verification=60, fit=60)
    ws = build(store, tmp_path)[SHEET_EXCLUDED]
    by_name = {r[headers(ws).index("الاسم")]: r for r in rows(ws)}
    assert set(by_name) == {"كبيرة", "ضعيفة"}
    assert "103" in by_name["كبيرة"][headers(ws).index("سبب الاستبعاد")]


def test_demo_leads_are_never_exported(store, tmp_path):
    seed_demo(store)
    lead(store, "حقيقي", path=UNDECIDED, verification=60, fit=60)
    wb = build(store, tmp_path)
    all_names = [r[1] for name in (SHEET_REVIEW, SHEET_EXCLUDED) for r in rows(wb[name])]
    assert all_names == ["حقيقي"]


def test_scores_are_numbers_and_contacts_are_filled(store, tmp_path):
    lead(store, "عميل", path=UNDECIDED, verification=63, fit=55, fields={
        "phone": ("+966541344479", "https://a.example", "own_site"),
        "email": ("info@a.com", "https://a.example", "own_site"),
        "decision_maker_name": ("د. فلان", "https://a.example", "own_site"),
    })
    ws = build(store, tmp_path)[SHEET_REVIEW]
    row = dict(zip(headers(ws), rows(ws)[0]))
    assert row["درجة التحقق"] == 63 and isinstance(row["درجة التحقق"], int)
    assert row["الهاتف"] == "+966541344479" and row["الإيميل"] == "info@a.com"
    assert row["صاحب القرار"] == "د. فلان"


def test_links_are_clickable_and_concerns_are_listed(store, tmp_path):
    lead(store, "عميل", path=UNDECIDED, verification=60, fit=60,
         fields={"website": ("https://a.example", "https://a.example", "own_site")},
         concerns=["تضارب هاتف", "نفس العنوان لكيان آخر"])
    ws = build(store, tmp_path)[SHEET_REVIEW]
    cols = headers(ws)
    cell = ws.cell(row=2, column=cols.index("الموقع") + 1)
    assert cell.hyperlink is not None and cell.hyperlink.target == "https://a.example"
    concerns = ws.cell(row=2, column=cols.index("تنبيهات للمراجع") + 1).value
    assert "تضارب هاتف" in concerns and "نفس العنوان" in concerns


def test_segment_and_country_are_shown_in_arabic(store, tmp_path):
    lead(store, "عميل", code="TEC-2", country="QA", city="الدوحة", path=UNDECIDED, verification=60, fit=60)
    ws = build(store, tmp_path)[SHEET_REVIEW]
    row = dict(zip(headers(ws), rows(ws)[0]))
    assert row["الدولة"] == "قطر" and "TEC-2" in row["الشريحة"] and "شركة" in row["الشريحة"]
    assert row["الحالة"] == "غير محسوم"


def test_sheets_are_right_to_left_with_frozen_header_and_filter(store, tmp_path):
    lead(store, "عميل", path=UNDECIDED, verification=60, fit=60)
    wb = build(store, tmp_path)
    for name in (SHEET_REVIEW, SHEET_EXCLUDED):
        ws = wb[name]
        assert ws.sheet_view.rightToLeft is True
        assert ws.freeze_panes == "C2"
        assert ws.auto_filter.ref
    assert wb[SHEET_REVIEW]["A1"].font.bold is True


def test_pending_lead_message_is_included_for_the_sales_manager(store, tmp_path):
    lead_id = lead(store, "عميل", path=(State.ENRICHED, State.VERIFIED, State.SCORED, State.DRAFTED, State.PENDING_APPROVAL),
                   verification=80, fit=75)
    store.set_field(lead_id, "message", "نص الرسالة المقترحة", "system", "system", 1.0)
    ws = build(store, tmp_path)[SHEET_REVIEW]
    row = dict(zip(headers(ws), rows(ws)[0]))
    assert row["الرسالة المقترحة"] == "نص الرسالة المقترحة"


def test_summary_counts_match_the_data(store, tmp_path):
    lead(store, "أ", path=UNDECIDED, verification=60, fit=60)
    lead(store, "ب", code="TEC-2", country="QA", path=OUT, verification=80)
    lead(store, "ج", path=DISCARDED, verification=20)
    ws = build(store, tmp_path)[SHEET_SUMMARY]
    text = " ".join(str(c.value) for r in ws.iter_rows() for c in r if c.value is not None)
    assert "3" in text                        # الإجمالي
    assert "لقطة" in text                      # تنويه أن الأرقام ثابتة


def test_notes_sheet_explains_scores_and_caveats(store, tmp_path):
    ws = build(store, tmp_path)[SHEET_NOTES]
    text = " ".join(str(c.value) for r in ws.iter_rows() for c in r if c.value)
    for needle in ("درجة التحقق", "درجة الملاءمة", "غير محسوم", "قبل التواصل", "مصادر عامة"):
        assert needle in text


def test_empty_database_still_produces_a_valid_workbook(store, tmp_path):
    wb = build(store, tmp_path)
    assert rows(wb[SHEET_REVIEW]) == []


def test_every_cell_uses_a_professional_font(store, tmp_path):
    lead(store, "عميل", path=UNDECIDED, verification=60, fit=60)
    wb = build(store, tmp_path)
    fonts = {c.font.name for ws in wb for r in ws.iter_rows() for c in r if c.value is not None}
    assert fonts == {"Arial"}


def test_channel_is_suggested_for_undecided_leads_from_available_contacts(store, tmp_path):
    lead(store, "مقهى", code="VIS-3", path=UNDECIDED, verification=60, fit=60,
         fields={"phone": ("0541344479", "https://a.example", "own_site")})
    lead(store, "شركة", code="TEC-2", country="QA", city="الدوحة", path=UNDECIDED, verification=60, fit=60,
         fields={"email": ("info@a.com", "https://a.example", "own_site")})
    ws = build(store, tmp_path)[SHEET_REVIEW]
    by_name = {r[1]: dict(zip(headers(ws), r)) for r in rows(ws)}
    assert by_name["مقهى"]["القناة المقترحة"] == "واتساب"
    assert by_name["شركة"]["القناة المقترحة"] == "إيميل"


def test_existing_channel_choice_is_not_overridden(store, tmp_path):
    lead(store, "عميل", code="VIS-3", path=UNDECIDED, verification=60, fit=60,
         fields={"phone": ("0541344479", "https://a.example", "own_site"),
                 "channel": ("email", "system", "system")})
    ws = build(store, tmp_path)[SHEET_REVIEW]
    assert dict(zip(headers(ws), rows(ws)[0]))["القناة المقترحة"] == "إيميل"
